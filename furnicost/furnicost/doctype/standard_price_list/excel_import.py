"""Read flat price tables and the two-sheet Phu Chau workbook. No database writes."""
import math
import re
import unicodedata
from collections import Counter
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook


def normalize(value):
	value = unicodedata.normalize("NFD", str(value or "").lower().replace("đ", "d"))
	return re.sub(r"[^a-z0-9]", "", "".join(c for c in value if not unicodedata.combining(c)))


ALIASES = {
	"item_code": ["Item Code", "Mã sản phẩm", "Mã item", "Mã hàng"],
	"item_name": ["Item Name", "Tên sản phẩm", "Tên item"],
	"category": ["category", "Hạng mục", "Nhóm sản phẩm"],
	"carcass_material": ["carcass_material", "Vật liệu thùng"],
	"door_material": ["door_material", "Vật liệu cánh"],
	"back_material": ["back_material", "Vật liệu hậu"],
	"description": ["description", "Mô tả", "Mô tả vật liệu"],
	"uom": ["UOM", "Đơn vị", "Đơn vị tính", "ĐVT"],
	"rate": ["rate", "Giá bán", "Đơn giá", "Đơn giá áp dụng"],
	"specification": ["Quy cách / ghi chú nguồn", "Quy cách"],
	"source_note": ["Diễn giải / cần xác nhận"],
}


def table(sheet, required):
	for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
		if number > 30:
			break
		headers = {normalize(cell): i for i, cell in enumerate(row) if cell is not None}
		mapping = {}
		for field, labels in ALIASES.items():
			for label in [field, *labels]:
				if normalize(label) in headers:
					mapping[field] = headers[normalize(label)]
					break
		if set(required).issubset(mapping):
			return [(idx, {key: values[col] if col < len(values) else None for key, col in mapping.items()})
				for idx, values in enumerate(sheet.iter_rows(min_row=number + 1, values_only=True), number + 1)]
	return None


def price(value):
	if value is None or str(value).strip() == "":
		raise ValueError("thiếu giá bán")
	if isinstance(value, str):
		value = re.sub(r"\s|VNĐ|VND|₫|đ", "", value, flags=re.I)
		if "," in value and "." in value:
			decimal, thousands = (",", ".") if value.rfind(",") > value.rfind(".") else (".", ",")
			value = value.replace(thousands, "").replace(decimal, ".")
		elif re.fullmatch(r"\d{1,3}([.,]\d{3})+", value):
			value = value.replace(".", "").replace(",", "")
		else:
			value = value.replace(",", ".")
	try:
		value = float(value)
	except (ValueError, TypeError):
		raise ValueError("giá bán không hợp lệ") from None
	if not math.isfinite(value) or value < 0:
		raise ValueError("giá bán phải là số không âm")
	return value


def parse_workbook(content):
	try:
		with ZipFile(BytesIO(content)) as archive:
			if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
				raise ValueError("Nội dung Excel vượt quá 50 MB.")
		workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
	except BadZipFile:
		raise ValueError("File không phải Excel .xlsx hợp lệ.") from None
	try:
		for sheet in workbook:
			rows = table(sheet, ["item_code", "item_name", "category", "rate"])
			if rows is not None:
				break
		else:
			raise ValueError("Không tìm thấy bảng có Item Code, Item Name, Hạng mục và Giá bán.")
		source_sheet = sheet.title
		item_codes = {str(row.get("item_code") or "").strip() for _, row in rows}
		refs = {}
		for sheet in workbook:
			if normalize(sheet.title) == normalize("Đối chiếu barem"):
				for _, row in table(sheet, ["item_code", "uom"]) or []:
					code = str(row.get("item_code") or "").strip()
					if code not in item_codes or not code:
						continue
					if code in refs:
						raise ValueError(f"Mã trùng trong sheet đối chiếu: {code}")
					if code:
						refs[code] = row
		counts = Counter(str(row.get("item_code") or "").strip().casefold() for _, row in rows)
		result = {"items": [], "warnings": [], "sheet": source_sheet}
		for number, row in rows:
			if not any(value is not None and str(value).strip() for value in row.values()):
				continue
			try:
				item = {field: str(row.get(field) if row.get(field) is not None else "").strip() for field in ALIASES if field not in ("rate", "specification", "source_note")}
				for field in ("item_code", "item_name", "category"):
					if not item[field]:
						raise ValueError(f"thiếu {field}")
				if counts[item["item_code"].casefold()] > 1:
					raise ValueError(f'mã bị trùng: {item["item_code"]}')
				item["rate"] = price(row.get("rate"))
				ref = refs.get(item["item_code"], {})
				if ref.get("rate") is not None and price(ref["rate"]) != item["rate"]:
					raise ValueError("giá khác sheet Đối chiếu barem")
				item["uom"] = item["uom"] or str(ref.get("uom") or "").strip()
				if not item["description"]:
					parts = [f"{label}: {item[field]}" for field, label in [("carcass_material", "Thùng"), ("door_material", "Cánh"), ("back_material", "Hậu")] if item[field]]
					for field, label in [("specification", "Quy cách"), ("source_note", "Ghi chú nguồn")]:
						value = row.get(field) or ref.get(field)
						if value:
							parts.append(f"{label}: {value}")
					item["description"] = "\n".join(parts)
				result["items"].append(item)
			except ValueError as error:
				result["warnings"].append(f"Dòng {number}: {error}")
		return result
	finally:
		workbook.close()


def export_workbook(rows):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	columns = [
		("item_code", "Item Code", 18), ("item_name", "Item Name", 55),
		("category", "Hạng mục", 25), ("carcass_material", "Vật liệu thùng", 30),
		("door_material", "Vật liệu cánh", 35), ("back_material", "Vật liệu hậu", 28),
		("description", "Mô tả", 65), ("uom", "Đơn vị", 12), ("rate", "Giá bán", 20),
	]
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Standard Price List"
	sheet.append([label for _, label, _ in columns])
	for column, (_, _, width) in enumerate(columns, 1):
		sheet.column_dimensions[get_column_letter(column)].width = width
		cell = sheet.cell(1, column)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="334155")
	for index, row in enumerate(rows, 2):
		for column, (field, _, _) in enumerate(columns, 1):
			value = row.get(field)
			cell = sheet.cell(index, column)
			if field == "rate":
				cell.value = price(value) if value is not None and str(value).strip() else None
				cell.number_format = "#,##0.##"
			else:
				cell.value = str(value) if value is not None else ""
				cell.data_type = "s"  # Preserve leading zeroes and literal text beginning with '='.
			cell.alignment = Alignment(vertical="top", wrap_text=True)
	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = sheet.dimensions
	stream = BytesIO()
	workbook.save(stream)
	workbook.close()
	return stream.getvalue()
